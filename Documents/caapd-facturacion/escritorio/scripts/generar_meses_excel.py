"""
scripts/generar_meses_excel.py — Limpia el mes pasado (2026-09) y genera el nuevo mes (2026-10).

Toma:
1. datos/RELACION_DE_USUARIOS_DEL_CAAPD.xlsx (lecturas reales del mes pasado).
2. Padrón maestro de usuarios en Firestore (DNI, categoría, sector, dirección oficial).

Produce:
1. datos/PADRON_MES_2026-09_LIMPIO.xlsx:
   Excel limpio con formato estándar de importación con lecturas anterior y actual del mes pasado.
2. datos/PADRON_MES_2026-10_NUEVO.xlsx:
   Excel para el nuevo mes donde la lectura anterior es la lectura actual de septiembre
   y la lectura actual queda en blanco lista para que el lector la tome.
"""
from __future__ import annotations
import os
import sys
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import codigos
import firebase_client as fb


def limpiar_y_generar_excels() -> None:
    ruta_relacion = BASE_DIR / "datos" / "RELACION_DE_USUARIOS_DEL_CAAPD.xlsx"
    if not ruta_relacion.exists():
        raise FileNotFoundError(f"No existe: {ruta_relacion}")

    print("Cargando padrón maestro desde Firestore...")
    usuarios_maestro = fb.cargar_usuarios(forzar=True)
    print(f"Total abonados maestros: {len(usuarios_maestro)}")

    # Indexar usuarios por medidor y por conexión
    u_por_medidor: Dict[str, Dict[str, Any]] = {}
    u_por_conexion: Dict[str, Dict[str, Any]] = {}
    for uid, u in usuarios_maestro.items():
        con = u.get("nro_conexion", "")
        if con:
            u_por_conexion[con] = u
        for m in u.get("medidores", []):
            m_norm = codigos.normalizar_medidor(m)
            u_por_medidor[m_norm] = u
            if m_norm.isdigit():
                u_por_medidor[m_norm.lstrip("0")] = u

    # Leer RELACION_DE_USUARIOS_DEL_CAAPD.xlsx
    wb_rel = openpyxl.load_workbook(str(ruta_relacion), data_only=True)

    registros_2026_09: List[Dict[str, Any]] = []
    medidores_procesados = set()

    for sheetname in wb_rel.sheetnames:
        ws = wb_rel[sheetname]
        filas = list(ws.iter_rows(values_only=True))
        if len(filas) < 2:
            continue

        for r in filas[1:]:
            if not any(r):
                continue

            raw_con = r[3] if len(r) > 3 else None
            raw_med = r[4] if len(r) > 4 else None
            raw_ben = r[5] if len(r) > 5 else ""
            raw_dir = r[6] if len(r) > 6 else ""
            raw_ant = r[7] if len(r) > 7 else 0
            raw_act = r[8] if len(r) > 8 else None

            con_norm = codigos.normalizar_conexion(raw_con) if raw_con else "000000"
            med_norm = codigos.normalizar_medidor(raw_med) if raw_med else ""

            # Vincular con usuario maestro
            u_match = None
            if med_norm and med_norm in u_por_medidor:
                u_match = u_por_medidor[med_norm]
            elif med_norm and med_norm.isdigit() and med_norm.lstrip("0") in u_por_medidor:
                u_match = u_por_medidor[med_norm.lstrip("0")]
            elif con_norm and con_norm in u_por_conexion:
                u_match = u_por_conexion[con_norm]

            if u_match:
                uid = str(u_match.get("_id") or u_match.get("id_usuario"))
                nombre = u_match.get("nombre", raw_ben)
                sector = u_match.get("sector", sheetname)
                direccion = u_match.get("direccion", raw_dir)
                categoria = u_match.get("codigo_categoria", "A")
            else:
                uid = f"SD-{int(con_norm):03d}" if con_norm.isdigit() else "SD-001"
                nombre = str(raw_ben or "").strip().upper()
                sector = sheetname
                direccion = str(raw_dir or "S/N").strip()
                categoria = "A"

            # Parsear lecturas
            try:
                lant_val = int(raw_ant) if raw_ant is not None and str(raw_ant).strip() != "" else 0
            except Exception:
                lant_val = 0

            try:
                lact_val = int(raw_act) if raw_act is not None and str(raw_act).strip() != "" else None
            except Exception:
                lact_val = None

            # Novedad y observaciones especiales
            novedad = "NORMAL"
            nota = ""
            if lact_val is not None:
                if lact_val < lant_val:
                    # Casos claros de cambio de medidor (nuevo medidor arranca en números bajos)
                    if (lant_val > 500 and lact_val < 50) or (lact_val - lant_val < -200):
                        novedad = "CAMBIO_MEDIDOR"
                        nota = f"Cambio de medidor: nuevo medidor inicia en {lact_val}"
                    else:
                        novedad = "NORMAL"
                        nota = f"Lectura regularizada (ant {lant_val} -> act {lact_val})"

            # Si hizo match con usuario maestro, usar su medidor canónico del padrón si coincide
            if u_match:
                meds_maestro = [codigos.normalizar_medidor(m) for m in u_match.get("medidores", [])]
                med_canonica = next(
                    (m for m in meds_maestro if m == med_norm or (m.isdigit() and med_norm.isdigit() and m.lstrip("0") == med_norm.lstrip("0"))),
                    med_norm,
                )
            else:
                med_canonica = med_norm

            reg = {
                "nro_conexion": con_norm,
                "dni_id": uid,
                "nombre": nombre,
                "sector": sector,
                "direccion": direccion,
                "categoria": categoria,
                "nro_medidor": med_canonica,
                "lectura_anterior": lant_val,
                "lectura_actual": lact_val,
                "novedad": novedad,
                "nota": nota,
            }
            registros_2026_09.append(reg)
            if med_canonica:
                medidores_procesados.add(med_canonica)

    # Solo agregar usuarios del padrón maestro que NO tengan ninguna fila en registros
    uids_procesados = {r["dni_id"] for r in registros_2026_09}
    conexiones_procesadas = {r["nro_conexion"] for r in registros_2026_09}

    for uid, u in usuarios_maestro.items():
        con = u.get("nro_conexion", "")
        if uid in uids_procesados or con in conexiones_procesadas:
            continue

        meds = [codigos.normalizar_medidor(m) for m in u.get("medidores", [])]
        m_unico = meds[0] if meds else ""
        registros_2026_09.append({
            "nro_conexion": con,
            "dni_id": uid,
            "nombre": u.get("nombre", ""),
            "sector": u.get("sector", "BARRIO DE CHAUPI"),
            "direccion": u.get("direccion", "S/N"),
            "categoria": u.get("codigo_categoria", "A"),
            "nro_medidor": m_unico,
            "lectura_anterior": 0,
            "lectura_actual": 0,
            "novedad": "NORMAL",
            "nota": "Abonado incorporado al padron",
        })

    # Ordenar por sector y conexión
    registros_2026_09.sort(key=lambda x: (x["sector"], int(x["nro_conexion"]) if x["nro_conexion"].isdigit() else 999999))

    print(f"Total registros procesados para 2026-09: {len(registros_2026_09)}")

    # -------------------------------------------------------------
    # 1. Guardar PADRON_MES_2026-09_LIMPIO.xlsx
    # -------------------------------------------------------------
    ruta_salida_09 = BASE_DIR / "datos" / "PADRON_MES_2026-09_LIMPIO.xlsx"
    _escribir_excel(
        ruta_salida_09,
        titulo_hoja="Lecturas Septiembre 2026",
        periodo_label="Septiembre 2026 (2026-09)",
        registros=registros_2026_09,
        es_mes_nuevo=False,
    )
    print(f"[OK] Creado: {ruta_salida_09}")

    # -------------------------------------------------------------
    # 2. Generar y guardar PADRON_MES_2026-10_NUEVO.xlsx
    # -------------------------------------------------------------
    registros_2026_10: List[Dict[str, Any]] = []
    for r in registros_2026_09:
        # En el nuevo mes, la lectura anterior es la lectura actual del mes pasado
        nueva_anterior = r["lectura_actual"] if r["lectura_actual"] is not None else r["lectura_anterior"]
        registros_2026_10.append({
            "nro_conexion": r["nro_conexion"],
            "dni_id": r["dni_id"],
            "nombre": r["nombre"],
            "sector": r["sector"],
            "direccion": r["direccion"],
            "categoria": r["categoria"],
            "nro_medidor": r["nro_medidor"],
            "lectura_anterior": nueva_anterior,
            "lectura_actual": None,  # Vacío para que el lector lo registre
            "novedad": "NORMAL",
            "nota": "",
        })

    ruta_salida_10 = BASE_DIR / "datos" / "PADRON_MES_2026-10_NUEVO.xlsx"
    _escribir_excel(
        ruta_salida_10,
        titulo_hoja="Lecturas Octubre 2026",
        periodo_label="Octubre 2026 (2026-10) — Nuevo Mes",
        registros=registros_2026_10,
        es_mes_nuevo=True,
    )
    print(f"[OK] Creado: {ruta_salida_10}")


def _escribir_excel(
    ruta_archivo: Path,
    titulo_hoja: str,
    periodo_label: str,
    registros: List[Dict[str, Any]],
    es_mes_nuevo: bool,
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lecturas"
    ws.views.sheetView[0].showGridLines = True

    # Encabezados
    headers = [
        "NRO_CONEXION",
        "DNI_ID",
        "NOMBRE",
        "SECTOR",
        "DIRECCION",
        "CATEGORIA",
        "NRO_MEDIDOR",
        "LECTURA_ANTERIOR",
        "LECTURA_ACTUAL",
        "NOVEDAD",
        "NOTA",
    ]

    fill_header = PatternFill(start_color="0284C7" if not es_mes_nuevo else "0D9488", end_color="0284C7" if not es_mes_nuevo else "0D9488", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for idx, reg in enumerate(registros, start=2):
        row_vals = [
            reg["nro_conexion"],
            reg["dni_id"],
            reg["nombre"],
            reg["sector"],
            reg["direccion"],
            reg["categoria"],
            reg["nro_medidor"],
            reg["lectura_anterior"],
            reg["lectura_actual"] if reg["lectura_actual"] is not None else "",
            reg["novedad"],
            reg["nota"],
        ]
        ws.append(row_vals)
        es_par = (idx % 2 == 0)

        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=idx, column=c_idx)
            cell.border = thin_border
            if es_par:
                cell.fill = fill_alt

            # Alineaciones específicas
            if c_idx in (1, 2, 6, 7, 10):
                cell.alignment = align_center
            elif c_idx in (8, 9):
                cell.alignment = align_right
            else:
                cell.alignment = align_left

        ws.row_dimensions[idx].height = 20

    # Auto-ajustar anchos
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            v_str = str(cell.value or "")
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws.column_dimensions["A"].width = 16  # CONEXION
    ws.column_dimensions["B"].width = 14  # DNI_ID
    ws.column_dimensions["C"].width = 38  # NOMBRE
    ws.column_dimensions["D"].width = 24  # SECTOR
    ws.column_dimensions["E"].width = 32  # DIRECCION
    ws.column_dimensions["G"].width = 16  # MEDIDOR

    wb.save(str(ruta_archivo))


if __name__ == "__main__":
    limpiar_y_generar_excels()
