"""
servicios/datos.py — Capa de servicio de datos para la aplicación de escritorio.

Proporciona una interfaz unificada para acceder a los datos, permitiendo
usar firebase_client en producción y dobles en memoria en pruebas unitarias.
Implementa importación perezosa (lazy) para que la ventana se abra al instante (< 1 s).
"""
import datetime
from typing import Any, Dict, List, Optional, Tuple
from servicios.estadisticas import calcular_estadisticas_inicio


def obtener_periodo_actual(fecha: Optional[datetime.date] = None) -> str:
    """Devuelve el periodo calendario actual en formato AAAA-MM."""
    return (fecha or datetime.date.today()).strftime("%Y-%m")


def siguiente_periodo(periodo: str) -> str:
    """Devuelve el mes siguiente a un periodo AAAA-MM."""
    fecha = datetime.datetime.strptime(periodo, "%Y-%m")
    if fecha.strftime("%Y-%m") != periodo:
        raise ValueError(f"Periodo inválido: {periodo}")
    anio = fecha.year + (1 if fecha.month == 12 else 0)
    mes = 1 if fecha.month == 12 else fecha.month + 1
    return f"{anio:04d}-{mes:02d}"


def todas_lecturas_completas(mes: Dict[str, Dict[str, Any]]) -> bool:
    """Indica si todos los medidores del mes tienen una lectura registrada."""
    medidores = [
        medidor
        for usuario in mes.values()
        if isinstance(usuario, dict)
        for medidor in (usuario.get("medidores") or [])
        if isinstance(medidor, dict)
    ]
    return bool(medidores) and all(
        str(medidor.get("estado_lectura") or "").upper() == "LEIDO"
        for medidor in medidores
    )


class ServicioDatos:
    """Interfaz base para el acceso a datos."""

    def cargar_config(self) -> Dict[str, Any]:
        raise NotImplementedError

    def guardar_config(self, datos: Dict[str, Any]) -> None:
        raise NotImplementedError

    def actualizar_config_campo(self, campo: str, valor: Any) -> None:
        raise NotImplementedError

    def cargar_usuarios(self, forzar: bool = False) -> Dict[str, Dict[str, Any]]:
        raise NotImplementedError

    def obtener_usuario(self, id_usuario: str) -> Optional[Dict[str, Any]]:
        raise NotImplementedError

    def guardar_usuario(
        self, id_usuario: str, datos: Dict[str, Any], agregar_al_mes: bool = False
    ) -> None:
        raise NotImplementedError

    def actualizar_usuario(self, id_usuario: str, cambios: Dict[str, Any]) -> None:
        raise NotImplementedError

    def corregir_dni_usuario(self, id_viejo: str, id_nuevo: str) -> Dict[str, Any]:
        raise NotImplementedError

    def obtener_siguiente_conexion_libre(self) -> str:
        raise NotImplementedError

    def obtener_siguiente_id_sd(self) -> str:
        raise NotImplementedError

    def validar_medidores_unicos(
        self, medidores: List[str], id_usuario_actual: Optional[str] = None
    ) -> Tuple[bool, Optional[str]]:
        raise NotImplementedError

    def cargar_mes(self, periodo: str) -> Dict[str, Dict[str, Any]]:
        raise NotImplementedError

    def verificar_puede_abrir_mes(self) -> Tuple[bool, str, str]:
        raise NotImplementedError

    def verificar_periodo_disponible(self, periodo: str) -> Tuple[bool, str]:
        raise NotImplementedError

    def abrir_mes(
        self,
        periodo: str,
        fecha_emision: str,
        fecha_vencimiento: str,
        lecturas_iniciales: Optional[Dict[str, int]] = None,
    ) -> Dict[str, Any]:
        raise NotImplementedError

    def cerrar_lecturas(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        raise NotImplementedError

    def reabrir_lecturas(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        raise NotImplementedError

    def guardar_lectura_medidor(
        self,
        periodo: str,
        id_usuario: str,
        nro_medidor: str,
        lectura_actual: int,
        novedad: str = "NORMAL",
        nota: str = "",
        registrado_por: str = "ADMIN",
    ) -> Dict[str, Any]:
        raise NotImplementedError

    def get_ops(self) -> Dict[str, int]:
        return {"lecturas": 0, "escrituras": 0}

    def hay_aviso_ops(self) -> Dict[str, Optional[str]]:
        return {"lecturas": None, "escrituras": None}

    def limpiar_cache(self, limpiar_config: bool = True) -> None:
        pass

    def listar_lectores(self) -> List[Dict[str, Any]]:
        raise NotImplementedError

    def crear_lector(self, email: str, password: str, display_name: Optional[str] = None) -> Dict[str, Any]:
        raise NotImplementedError

    def actualizar_lector(
        self,
        uid: str,
        display_name: Optional[str] = None,
        password: Optional[str] = None,
        disabled: Optional[bool] = None,
    ) -> Dict[str, Any]:
        raise NotImplementedError

    def eliminar_lector(self, uid: str) -> None:
        raise NotImplementedError

    def verificar_estado_generacion_recibos(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        raise NotImplementedError

    def estimar_lectura_medidor(
        self, periodo: str, id_usuario: str, nro_medidor: str, consumo_estimado: Optional[int] = None
    ) -> Dict[str, Any]:
        raise NotImplementedError

    def estimar_todas_lecturas_faltantes(self, periodo: str) -> Dict[str, Any]:
        raise NotImplementedError

    def generar_recibos_mes(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        raise NotImplementedError

    def generar_recibo_usuario(self, periodo: str, id_usuario: str) -> Dict[str, Any]:
        raise NotImplementedError

    def generar_pdf_recibo(self, periodo: str, id_usuario: str) -> str:
        raise NotImplementedError

    def generar_zip_recibos(self, periodo: str, sector: Optional[str] = None) -> Optional[str]:
        raise NotImplementedError

    def registrar_pago(
        self,
        periodo: str,
        id_usuario: str,
        monto_pagado: float,
        fecha_pago: Optional[str] = None,
    ) -> Dict[str, Any]:
        raise NotImplementedError

    def anular_pago(self, periodo: str, id_usuario: str) -> Dict[str, Any]:
        raise NotImplementedError

    def exportar_reporte_morosos_excel(
        self,
        periodo: Optional[str] = None,
        ruta_salida: Optional[str] = None,
    ) -> str:
        raise NotImplementedError

    def obtener_estadisticas_inicio(self) -> Tuple[Dict[str, Any], list, dict]:
        """
        Carga config, usuarios y mes_actual y calcula estadísticas.
        Retorna (config, stats_sectores, stats_recibos).
        """
        config = self.cargar_config()
        mes_actual = config.get("mes_actual", {})
        periodo = mes_actual.get("periodo")

        if not periodo:
            return config, [], {"generados": 0, "pagados": 0, "pendientes": 0}

        usuarios = self.cargar_usuarios()
        mes = self.cargar_mes(periodo)
        stats_sectores, stats_recibos = calcular_estadisticas_inicio(config, usuarios, mes)
        return config, stats_sectores, stats_recibos


class ServicioDatosFirebase(ServicioDatos):
    """Implementación real conectada a Firebase con carga diferida (lazy)."""

    def __init__(self) -> None:
        self._fb = None

    def _get_fb(self) -> Any:
        """Inicializa firebase_client en el primer uso (dentro del hilo de trabajo)."""
        if self._fb is None:
            import firebase_client as fb
            self._fb = fb
        return self._fb

    def cargar_config(self) -> Dict[str, Any]:
        return self._get_fb().cargar_config()

    def guardar_config(self, datos: Dict[str, Any]) -> None:
        from respaldo import on_datos_cambiados

        self._get_fb().guardar_config(datos)
        on_datos_cambiados("config_guardada")

    def actualizar_config_campo(self, campo: str, valor: Any) -> None:
        from respaldo import on_datos_cambiados

        self._get_fb().actualizar_config_campo(campo, valor)
        on_datos_cambiados("config_actualizada", {"campo": campo})

    def cargar_usuarios(self, forzar: bool = False) -> Dict[str, Dict[str, Any]]:
        return self._get_fb().cargar_usuarios(forzar=forzar)

    def obtener_usuario(self, id_usuario: str) -> Optional[Dict[str, Any]]:
        return self._get_fb().obtener_usuario(id_usuario)

    def guardar_usuario(
        self, id_usuario: str, datos: Dict[str, Any], agregar_al_mes: bool = False
    ) -> None:
        from respaldo import on_datos_cambiados

        fb = self._get_fb()
        fb.guardar_usuario(id_usuario, datos)
        if agregar_al_mes:
            fb.agregar_usuario_al_mes_actual(id_usuario, datos)
        on_datos_cambiados("usuario_guardado", {"id_usuario": id_usuario})

    def actualizar_usuario(self, id_usuario: str, cambios: Dict[str, Any]) -> None:
        from respaldo import on_datos_cambiados

        fb = self._get_fb()
        fb.actualizar_usuario(id_usuario, cambios)
        on_datos_cambiados("usuario_estado", {"id_usuario": id_usuario, **cambios})

    def corregir_dni_usuario(self, id_viejo: str, id_nuevo: str) -> Dict[str, Any]:
        from respaldo import on_datos_cambiados

        fb = self._get_fb()
        res = fb.corregir_dni_usuario(id_viejo, id_nuevo)
        on_datos_cambiados("corregir_dni", res)
        return res

    def obtener_siguiente_conexion_libre(self) -> str:
        return self._get_fb().obtener_siguiente_conexion_libre()

    def obtener_siguiente_id_sd(self) -> str:
        return self._get_fb().obtener_siguiente_id_sd()

    def validar_medidores_unicos(
        self, medidores: List[str], id_usuario_actual: Optional[str] = None
    ) -> Tuple[bool, Optional[str]]:
        return self._get_fb().validar_medidores_unicos(
            medidores, id_usuario_actual=id_usuario_actual
        )

    def cargar_mes(self, periodo: str) -> Dict[str, Dict[str, Any]]:
        return self._get_fb().cargar_mes(periodo)

    def verificar_puede_abrir_mes(self) -> Tuple[bool, str, str]:
        return self._get_fb().verificar_puede_abrir_mes()

    def verificar_periodo_disponible(self, periodo: str) -> Tuple[bool, str]:
        return self._get_fb().verificar_periodo_disponible(periodo)

    def abrir_mes(
        self,
        periodo: str,
        fecha_emision: str,
        fecha_vencimiento: str,
        lecturas_iniciales: Optional[Dict[str, int]] = None,
    ) -> Dict[str, Any]:
        return self._get_fb().abrir_mes(
            periodo,
            fecha_emision,
            fecha_vencimiento,
            lecturas_iniciales=lecturas_iniciales,
        )

    def cerrar_lecturas(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        return self._get_fb().cerrar_lecturas(periodo)

    def reabrir_lecturas(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        return self._get_fb().reabrir_lecturas(periodo)

    def guardar_lectura_medidor(
        self,
        periodo: str,
        id_usuario: str,
        nro_medidor: str,
        lectura_actual: int,
        novedad: str = "NORMAL",
        nota: str = "",
        registrado_por: str = "ADMIN",
    ) -> Dict[str, Any]:
        return self._get_fb().guardar_lectura_medidor(
            periodo,
            id_usuario,
            nro_medidor,
            lectura_actual,
            novedad=novedad,
            nota=nota,
            registrado_por=registrado_por,
        )

    def get_ops(self) -> Dict[str, int]:
        if self._fb is None:
            return {"lecturas": 0, "escrituras": 0}
        return self._fb.get_ops()

    def hay_aviso_ops(self) -> Dict[str, Optional[str]]:
        if self._fb is None:
            return {"lecturas": None, "escrituras": None}
        return self._fb.hay_aviso_ops()

    def limpiar_cache(self, limpiar_config: bool = True) -> None:
        if self._fb is not None:
            self._fb.limpiar_cache(limpiar_config=limpiar_config)

    def _asegurar_firebase(self) -> None:
        """Asegura que Firebase Admin esté inicializado sin forzar la conexión gRPC de Firestore."""
        fb = self._get_fb()
        if getattr(fb, "_app", None) is None:
            fb._inicializar()

    def listar_lectores(self) -> List[Dict[str, Any]]:
        self._asegurar_firebase()
        from firebase_admin import auth
        page = auth.list_users()
        lectores = []
        for u in page.users:
            lectores.append({
                "uid": u.uid,
                "email": u.email or "",
                "display_name": u.display_name or "",
                "disabled": bool(u.disabled),
                "created_at": getattr(u.user_metadata, "creation_timestamp", None),
                "last_sign_in": getattr(u.user_metadata, "last_sign_in_timestamp", None),
            })
        return lectores

    def crear_lector(self, email: str, password: str, display_name: Optional[str] = None) -> Dict[str, Any]:
        self._asegurar_firebase()
        from firebase_admin import auth
        user = auth.create_user(
            email=email.strip().lower(),
            password=password.strip(),
            display_name=display_name.strip() if display_name else None,
        )
        return {
            "uid": user.uid,
            "email": user.email,
            "display_name": user.display_name or "",
            "disabled": bool(user.disabled),
        }

    def actualizar_lector(
        self,
        uid: str,
        display_name: Optional[str] = None,
        password: Optional[str] = None,
        disabled: Optional[bool] = None,
    ) -> Dict[str, Any]:
        self._asegurar_firebase()
        from firebase_admin import auth
        kwargs: Dict[str, Any] = {}
        if display_name is not None:
            kwargs["display_name"] = display_name.strip()
        if password is not None and password.strip():
            kwargs["password"] = password.strip()
        if disabled is not None:
            kwargs["disabled"] = disabled
        user = auth.update_user(uid, **kwargs)
        return {
            "uid": user.uid,
            "email": user.email,
            "display_name": user.display_name or "",
            "disabled": bool(user.disabled),
        }

    def eliminar_lector(self, uid: str) -> None:
        self._asegurar_firebase()
        from firebase_admin import auth
        auth.delete_user(uid)

    def verificar_estado_generacion_recibos(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        return self._get_fb().verificar_estado_generacion_recibos(periodo)

    def estimar_lectura_medidor(
        self, periodo: str, id_usuario: str, nro_medidor: str, consumo_estimado: Optional[int] = None
    ) -> Dict[str, Any]:
        return self._get_fb().estimar_lectura_medidor(periodo, id_usuario, nro_medidor, consumo_estimado)

    def estimar_todas_lecturas_faltantes(self, periodo: str) -> Dict[str, Any]:
        return self._get_fb().estimar_todas_lecturas_faltantes(periodo)

    def generar_recibos_mes(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        return self._get_fb().generar_recibos_mes(periodo)

    def generar_recibo_usuario(self, periodo: str, id_usuario: str) -> Dict[str, Any]:
        return self._get_fb().generar_recibo_usuario(periodo, id_usuario)

    def generar_pdf_recibo(self, periodo: str, id_usuario: str) -> str:
        cfg = self.cargar_config()
        doc = self.cargar_mes(periodo).get(id_usuario)
        if not doc:
            raise ValueError(f"No se encontró el recibo para el usuario {id_usuario} en el mes {periodo}.")
        from servicios.generador_pdf import generar_recibo_pdf_para_usuario
        ruta = generar_recibo_pdf_para_usuario(doc, cfg, periodo)
        return str(ruta)

    def generar_zip_recibos(self, periodo: str, sector: Optional[str] = None) -> Optional[str]:
        cfg = self.cargar_config()
        mes_docs = self.cargar_mes(periodo)
        from servicios.generador_pdf import generar_recibo_pdf_para_usuario, empaquetar_recibos_en_zip
        for uid, doc in mes_docs.items():
            if sector and doc.get("sector") != sector:
                continue
            if doc.get("nro_recibo"):
                generar_recibo_pdf_para_usuario(doc, cfg, periodo)
        ruta_zip = empaquetar_recibos_en_zip(periodo, sector)
        return str(ruta_zip) if ruta_zip else None

    def registrar_pago(
        self,
        periodo: str,
        id_usuario: str,
        monto_pagado: float,
        fecha_pago: Optional[str] = None,
    ) -> Dict[str, Any]:
        return self._get_fb().registrar_pago(
            periodo, id_usuario, monto_pagado, fecha_pago=fecha_pago
        )

    def anular_pago(self, periodo: str, id_usuario: str) -> Dict[str, Any]:
        return self._get_fb().anular_pago(periodo, id_usuario)

    def exportar_reporte_morosos_excel(
        self,
        periodo: Optional[str] = None,
        ruta_salida: Optional[str] = None,
    ) -> str:
        return self._get_fb().exportar_reporte_morosos_excel(
            periodo=periodo, ruta_salida=ruta_salida
        )


class ServicioDatosMemoria(ServicioDatos):
    """Doble en memoria para pruebas sin llamadas externas ni red."""

    def __init__(
        self,
        config: Optional[Dict[str, Any]] = None,
        usuarios: Optional[Dict[str, Dict[str, Any]]] = None,
        mes: Optional[Dict[str, Dict[str, Any]]] = None,
        ops: Optional[Dict[str, int]] = None,
    ) -> None:
        self.config = config or {
            "sectores": [
                "BARRIO DE CHAUPI", "BARRIO DE PACHAYZA", "BARRIO DE LUCMANI",
                "HUAQUILLO", "CHAUPICERCA", "AYACUCHO"
            ],
            "categorias": {
                "A": {"nombre": "Doméstico"},
                "B": {"nombre": "Comercial"},
            },
            "mes_actual": {"estado": "ABIERTO", "periodo": "2026-03"},
        }
        self.usuarios = usuarios or {}
        self.mes = mes or {}
        self.ops = ops or {"lecturas": 5, "escrituras": 2}
        self.limpiezas_cache = 0
        self.lectores: Dict[str, Dict[str, Any]] = {}
        self.es_memoria: bool = True

    def cargar_config(self) -> Dict[str, Any]:
        return self.config

    def guardar_config(self, datos: Dict[str, Any]) -> None:
        self.config = dict(datos)

    def actualizar_config_campo(self, campo: str, valor: Any) -> None:
        partes = campo.split(".")
        nodo = self.config
        for p in partes[:-1]:
            nodo = nodo.setdefault(p, {})
        nodo[partes[-1]] = valor

    def cargar_usuarios(self, forzar: bool = False) -> Dict[str, Dict[str, Any]]:
        return self.usuarios

    def obtener_usuario(self, id_usuario: str) -> Optional[Dict[str, Any]]:
        return self.usuarios.get(id_usuario)

    def guardar_usuario(
        self, id_usuario: str, datos: Dict[str, Any], agregar_al_mes: bool = False
    ) -> None:
        d = dict(datos)
        d["_id"] = id_usuario
        self.usuarios[id_usuario] = d
        if agregar_al_mes and self.config.get("mes_actual", {}).get("estado") == "ABIERTO":
            periodo = self.config.get("mes_actual", {}).get("periodo", "2026-03")
            if id_usuario not in self.mes:
                self.mes[id_usuario] = {
                    "nro_conexion": d.get("nro_conexion", ""),
                    "nombre": d.get("nombre", ""),
                    "sector": d.get("sector", ""),
                    "direccion": d.get("direccion", ""),
                    "codigo_categoria": d.get("codigo_categoria", "A"),
                    "medidores": [
                        {"nro_medidor": m, "lectura_anterior": 0, "lectura_actual": None}
                        for m in d.get("medidores", [])
                    ],
                }

    def actualizar_usuario(self, id_usuario: str, cambios: Dict[str, Any]) -> None:
        if id_usuario in self.usuarios:
            self.usuarios[id_usuario].update(cambios)

    def corregir_dni_usuario(self, id_viejo: str, id_nuevo: str) -> Dict[str, Any]:
        if id_viejo == id_nuevo:
            raise ValueError("El nuevo ID es igual al ID actual.")
        if id_nuevo in self.usuarios:
            raise ValueError(f"El usuario con ID {id_nuevo} ya existe en el padrón.")
        if id_viejo not in self.usuarios:
            raise ValueError(f"El usuario con ID {id_viejo} no existe.")

        datos = self.usuarios.pop(id_viejo)
        datos["_id"] = id_nuevo
        self.usuarios[id_nuevo] = datos

        meses_act = []
        if id_viejo in self.mes:
            doc_mes = self.mes.pop(id_viejo)
            self.mes[id_nuevo] = doc_mes
            meses_act.append(self.config.get("mes_actual", {}).get("periodo", "2026-03"))

        return {
            "id_viejo": id_viejo,
            "id_nuevo": id_nuevo,
            "nombre": datos.get("nombre", ""),
            "usuario_actualizado": True,
            "meses_actualizados": meses_act,
        }

    def obtener_siguiente_conexion_libre(self) -> str:
        max_num = 0
        for u in self.usuarios.values():
            con = str(u.get("nro_conexion") or "").strip()
            if con.isdigit():
                max_num = max(max_num, int(con))
        return f"{max_num + 1:06d}"

    def obtener_siguiente_id_sd(self) -> str:
        import re
        max_sd = 0
        for uid in self.usuarios.keys():
            m = re.match(r"^SD-(\d+)$", str(uid).strip())
            if m:
                max_sd = max(max_sd, int(m.group(1)))
        return f"SD-{max_sd + 1:03d}"

    def validar_medidores_unicos(
        self, medidores: List[str], id_usuario_actual: Optional[str] = None
    ) -> Tuple[bool, Optional[str]]:
        meds_set = {m.strip() for m in medidores if m and m.strip()}
        for uid, u in self.usuarios.items():
            if id_usuario_actual and uid == id_usuario_actual:
                continue
            u_meds = [
                m if isinstance(m, str) else m.get("nro_medidor", "")
                for m in u.get("medidores", [])
            ]
            u_meds_set = {m.strip() for m in u_meds if m and m.strip()}
            interseccion = meds_set.intersection(u_meds_set)
            if interseccion:
                med_duplicado = list(interseccion)[0]
                nombre = u.get("nombre", "Desconocido")
                return (
                    False,
                    f"El medidor {med_duplicado} ya está asignado a '{nombre}' (ID: {uid}).",
                )
        return True, None

    def cargar_mes(self, periodo: str) -> Dict[str, Dict[str, Any]]:
        return self.mes

    def verificar_puede_abrir_mes(self) -> Tuple[bool, str, str]:
        mes_act = self.config.get("mes_actual") or {}
        periodo_actual = mes_act.get("periodo")
        estado_actual = (mes_act.get("estado") or "").upper()

        periodo_calendario = obtener_periodo_actual()
        siguiente = siguiente_periodo(periodo_actual) if periodo_actual else periodo_calendario

        if periodo_actual and estado_actual == "ABIERTO":
            return (
                False,
                f"Ya existe una base activa para el mes {periodo_actual} (Estado: {estado_actual}). "
                f"Debe cerrarse y facturarse antes de abrir un nuevo mes.",
                periodo_actual,
            )

        lecturas_completas = todas_lecturas_completas(self.mes)
        if periodo_actual and estado_actual != "FACTURADO" and not (
            estado_actual == "CERRADO"
            and lecturas_completas
        ):
            return (
                False,
                f"Ya existe una base activa para el mes {periodo_actual} "
                f"(Estado: {estado_actual or 'SIN ESTADO'}). "
                "Debe cerrarse y facturarse antes de abrir un nuevo mes.",
                periodo_actual,
            )

        periodo_sugerido = max(siguiente, periodo_calendario)
        return True, "Listo para abrir un nuevo mes.", periodo_sugerido

    def verificar_periodo_disponible(self, periodo: str) -> Tuple[bool, str]:
        if hasattr(self, "meses_existentes") and periodo in self.meses_existentes:
            return (
                False,
                f"Ya existe una base de datos registrada para el mes {periodo}. "
                f"No es posible sobrescribir una base existente.",
            )
        # Si el mes actual ya tiene ese periodo y tiene datos
        if self.config.get("mes_actual", {}).get("periodo") == periodo and self.mes:
            return (
                False,
                f"Ya existe una base de datos registrada para el mes {periodo}. "
                f"No es posible sobrescribir una base existente.",
            )
        return True, f"El periodo {periodo} está disponible."

    def abrir_mes(
        self,
        periodo: str,
        fecha_emision: str,
        fecha_vencimiento: str,
        lecturas_iniciales: Optional[Dict[str, int]] = None,
    ) -> Dict[str, Any]:
        puede, motivo, _ = self.verificar_puede_abrir_mes()
        if not puede:
            raise ValueError(motivo)
        disp, mot_disp = self.verificar_periodo_disponible(periodo)
        if not disp:
            raise ValueError(mot_disp)

        activos = {uid: u for uid, u in self.usuarios.items() if u.get("activo", True)}
        if not activos:
            raise ValueError("No se encontraron usuarios activos en el padrón para abrir el mes.")

        lecturas_iniciales = lecturas_iniciales or {}
        nuevos_docs = {}
        for uid, u in activos.items():
            meds_mes = []
            for m in u.get("medidores", []):
                nro = m if isinstance(m, str) else m.get("nro_medidor", "")
                if nro:
                    meds_mes.append(
                        {
                            "nro_medidor": nro,
                            "lectura_anterior": int(lecturas_iniciales.get(nro, 0)),
                            "fecha_anterior": None,
                            "consumo_anterior": None,
                            "lectura_actual": None,
                            "fecha_lectura": None,
                            "estado_lectura": "PENDIENTE",
                            "novedad": "",
                            "nota": "",
                            "registrado_por": None,
                        }
                    )
            nuevos_docs[uid] = {
                "nro_conexion": u.get("nro_conexion", ""),
                "nombre": u.get("nombre", ""),
                "sector": u.get("sector", ""),
                "direccion": u.get("direccion", ""),
                "codigo_categoria": u.get("codigo_categoria", "A"),
                "medidores": meds_mes,
                "consumo_total_m3": None,
                "historial_m3": [],
                "nro_recibo": None,
                "monto_agua": None,
                "monto_desague": None,
                "cargo_fijo": None,
                "deuda_anterior": 0.0,
                "igv": 0.0,
                "redondeo": 0.0,
                "total_a_pagar": None,
                "meses_adeudados": 1,
                "monto_pagado": 0.0,
                "fecha_pago": None,
            }

        self.mes = nuevos_docs
        self.config["mes_actual"] = {
            "periodo": periodo,
            "fecha_emision": fecha_emision,
            "fecha_vencimiento": fecha_vencimiento,
            "estado": "ABIERTO",
        }
        return {"periodo": periodo, "usuarios_creados": len(nuevos_docs), "estado": "ABIERTO"}

    def cerrar_lecturas(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        mes_act = self.config.get("mes_actual") or {}
        p = periodo or mes_act.get("periodo")
        if not p:
            raise ValueError("No hay un mes activo para cerrar lecturas.")
        self.config["mes_actual"]["estado"] = "CERRADO"
        return {"periodo": p, "estado": "CERRADO"}

    def reabrir_lecturas(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        mes_act = self.config.get("mes_actual") or {}
        p = periodo or mes_act.get("periodo")
        if not p:
            raise ValueError("No hay un mes activo para reabrir.")
        if mes_act.get("estado") == "FACTURADO":
            raise ValueError("El mes ya fue facturado y no puede reabrirse para lecturas.")
        self.config["mes_actual"]["estado"] = "ABIERTO"
        return {"periodo": p, "estado": "ABIERTO"}

    def guardar_lectura_medidor(
        self,
        periodo: str,
        id_usuario: str,
        nro_medidor: str,
        lectura_actual: int,
        novedad: str = "NORMAL",
        nota: str = "",
        registrado_por: str = "ADMIN",
    ) -> Dict[str, Any]:
        import datetime
        import calculos

        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError(f"Usuario {id_usuario} no encontrado en el mes {periodo}.")

        med_obj = next((m for m in doc.get("medidores", []) if m.get("nro_medidor") == nro_medidor), None)
        if not med_obj:
            raise ValueError(f"Medidor {nro_medidor} no pertenece a {id_usuario}.")

        med_obj["lectura_actual"] = int(lectura_actual)
        med_obj["fecha_lectura"] = datetime.date.today().isoformat()
        med_obj["estado_lectura"] = "LEIDO"
        med_obj["novedad"] = novedad or "NORMAL"
        med_obj["nota"] = nota or ""
        med_obj["registrado_por"] = registrado_por

        meds = doc.get("medidores", [])
        if all(m.get("lectura_actual") is not None for m in meds):
            total_c = 0.0
            for m in meds:
                c = calculos.consumo_medidor(
                    m.get("lectura_anterior") or 0,
                    m["lectura_actual"],
                    m.get("novedad"),
                )
                total_c += c
            doc["consumo_total_m3"] = round(total_c, 2)
        else:
            doc["consumo_total_m3"] = None

        return {"id_usuario": id_usuario, "nro_medidor": nro_medidor, "consumo_total_m3": doc.get("consumo_total_m3")}

    def get_ops(self) -> Dict[str, int]:
        return self.ops

    def hay_aviso_ops(self) -> Dict[str, Optional[str]]:
        return {"lecturas": None, "escrituras": None}

    def limpiar_cache(self, limpiar_config: bool = True) -> None:
        self.limpiezas_cache += 1

    def listar_lectores(self) -> List[Dict[str, Any]]:
        return list(self.lectores.values())

    def crear_lector(self, email: str, password: str, display_name: Optional[str] = None) -> Dict[str, Any]:
        uid = f"uid_{email.strip().lower()}"
        lector = {
            "uid": uid,
            "email": email.strip().lower(),
            "display_name": display_name.strip() if display_name else "",
            "disabled": False,
            "created_at": 1700000000000,
            "last_sign_in": None,
        }
        self.lectores[uid] = lector
        return lector

    def actualizar_lector(
        self,
        uid: str,
        display_name: Optional[str] = None,
        password: Optional[str] = None,
        disabled: Optional[bool] = None,
    ) -> Dict[str, Any]:
        if uid not in self.lectores:
            raise ValueError(f"Lector con UID {uid} no encontrado.")
        if display_name is not None:
            self.lectores[uid]["display_name"] = display_name.strip()
        if disabled is not None:
            self.lectores[uid]["disabled"] = disabled
        return self.lectores[uid]

    def eliminar_lector(self, uid: str) -> None:
        if uid in self.lectores:
            del self.lectores[uid]

    def verificar_estado_generacion_recibos(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        mes_act = self.config.get("mes_actual") or {}
        periodo_obj = periodo or mes_act.get("periodo") or "2026-10"
        estado_mes = (mes_act.get("estado") or "ABIERTO").upper()
        docs = self.mes

        faltantes = []
        recibos_generados = 0
        total_facturado = 0.0
        pagados = 0
        pendientes = 0

        for uid, doc in docs.items():
            for m in doc.get("medidores", []):
                act = m.get("lectura_actual")
                st = (m.get("estado_lectura") or "").upper()
                if act is None or st in ("PENDIENTE", "SIN_ACCESO"):
                    faltantes.append({
                        "id_usuario": uid,
                        "nro_conexion": doc.get("nro_conexion", ""),
                        "nombre": doc.get("nombre", ""),
                        "sector": doc.get("sector", ""),
                        "nro_medidor": m.get("nro_medidor", ""),
                        "lectura_anterior": m.get("lectura_anterior", 0),
                        "consumo_anterior": m.get("consumo_anterior"),
                        "estado_lectura": st or "PENDIENTE",
                        "novedad": m.get("novedad", ""),
                    })
            if doc.get("nro_recibo"):
                recibos_generados += 1
                tot = float(doc.get("total_a_pagar") or 0.0)
                total_facturado += tot
                mp = float(doc.get("monto_pagado") or 0.0)
                if mp >= tot and tot > 0:
                    pagados += 1
                else:
                    pendientes += 1

        return {
            "periodo": periodo_obj,
            "estado": estado_mes,
            "puede_generar": len(faltantes) == 0,
            "faltantes": faltantes,
            "total_usuarios": len(docs),
            "recibos_generados": recibos_generados,
            "total_facturado": round(total_facturado, 2),
            "pagados": pagados,
            "pendientes": pendientes,
        }

    def estimar_lectura_medidor(
        self, periodo: str, id_usuario: str, nro_medidor: str, consumo_estimado: Optional[int] = None
    ) -> Dict[str, Any]:
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError(f"Usuario {id_usuario} no encontrado.")
        med_obj = next((m for m in doc.get("medidores", []) if m.get("nro_medidor") == nro_medidor), None)
        if not med_obj:
            raise ValueError(f"Medidor {nro_medidor} no encontrado.")
        if consumo_estimado is not None:
            c = int(consumo_estimado)
        else:
            c_ant = med_obj.get("consumo_anterior")
            if c_ant is None:
                raise ValueError("Sin consumo anterior.")
            c = int(c_ant)
        l_act = int(med_obj.get("lectura_anterior") or 0) + max(0, c)
        return self.guardar_lectura_medidor(
            periodo, id_usuario, nro_medidor, l_act, novedad="LECTURA_ESTIMADA", nota="Estimada por la secretaria", registrado_por="ADMIN"
        )

    def estimar_todas_lecturas_faltantes(self, periodo: str) -> Dict[str, Any]:
        est = self.verificar_estado_generacion_recibos(periodo)
        cnt = 0
        req = []
        for f in est.get("faltantes", []):
            if f.get("consumo_anterior") is not None:
                self.estimar_lectura_medidor(periodo, f["id_usuario"], f["nro_medidor"], int(f["consumo_anterior"]))
                cnt += 1
            else:
                req.append(f)
        return {"estimados": cnt, "requieren_manual": req}

    def generar_recibos_mes(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        import calculos
        periodo_obj = periodo or self.config.get("mes_actual", {}).get("periodo") or "2026-10"
        est = self.verificar_estado_generacion_recibos(periodo_obj)
        if not est["puede_generar"]:
            raise ValueError("Existen medidores con lecturas pendientes.")

        ultimo_recibo_num = int(self.config.get("ultimo_recibo") or 0)
        categorias = self.config.get("categorias", {})
        igv_cfg = self.config.get("igv") or {}
        igv_activo = bool(igv_cfg.get("activo", False))
        igv_porcentaje = float(igv_cfg.get("porcentaje", 0.18))
        sectores = self.config.get("sectores") or []

        def orden_k(item):
            uid, d = item
            sec = str(d.get("sector") or "")
            s_idx = sectores.index(sec) if sec in sectores else 999
            return (s_idx, sec, str(d.get("nombre") or "").upper())

        items_ord = sorted(self.mes.items(), key=orden_k)
        nuevos = 0
        tot_monto = 0.0

        for uid, doc in items_ord:
            if doc.get("nro_recibo"):
                n_rec = doc["nro_recibo"]
            else:
                ultimo_recibo_num += 1
                n_rec = f"{ultimo_recibo_num:09d}"
                nuevos += 1
            doc["nro_recibo"] = n_rec

            cat_cod = doc.get("codigo_categoria", "A")
            cat_obj = categorias.get(cat_cod) or categorias.get("A", {
                "tarifa_desague": 0.0,
                "cargo_fijo": 1.5,
                "modo_rangos": "unico",
                "rangos": [{"m3_desde": 0, "m3_hasta": None, "precio_m3": 1.0}],
            })

            calc = calculos.calcular_recibo(
                medidores=doc.get("medidores", []),
                categoria=cat_obj,
                igv_activo=igv_activo,
                igv_porcentaje=igv_porcentaje,
                recibo_anterior=None,
            )

            doc.update({
                "consumo_total_m3": calc["consumo_total_m3"],
                "monto_agua": calc["monto_agua"],
                "monto_desague": calc["monto_desague"],
                "cargo_fijo": calc["cargo_fijo"],
                "deuda_anterior": calc["deuda_anterior"],
                "igv": calc["igv"],
                "redondeo": calc["redondeo"],
                "total_a_pagar": calc["total_a_pagar"],
                "meses_adeudados": calc["meses_adeudados"],
                "monto_pagado": 0.0,
                "fecha_pago": None,
            })
            tot_monto += calc["total_a_pagar"]

        self.config["ultimo_recibo"] = ultimo_recibo_num
        self.config.setdefault("mes_actual", {})["estado"] = "FACTURADO"

        return {
            "periodo": periodo_obj,
            "total_recibos": len(self.mes),
            "nuevos_recibos": nuevos,
            "ultimo_recibo": ultimo_recibo_num,
            "total_facturado": round(tot_monto, 2),
        }

    def generar_recibo_usuario(self, periodo: str, id_usuario: str) -> Dict[str, Any]:
        import calculos
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError(f"Usuario {id_usuario} no encontrado en el mes {periodo}.")

        for m in doc.get("medidores", []):
            if m.get("lectura_actual") is None:
                m["lectura_actual"] = m.get("lectura_anterior") or 0
                m["estado_lectura"] = "ESTIMADO"

        ultimo_recibo_num = int(self.config.get("ultimo_recibo") or 0)
        if not doc.get("nro_recibo"):
            ultimo_recibo_num += 1
            doc["nro_recibo"] = f"{ultimo_recibo_num:09d}"
            self.config["ultimo_recibo"] = ultimo_recibo_num

        categorias = self.config.get("categorias", {})
        cat_cod = doc.get("codigo_categoria", "A")
        cat_obj = categorias.get(cat_cod) or categorias.get("A", {
            "tarifa_desague": 0.0,
            "cargo_fijo": 1.5,
            "modo_rangos": "unico",
            "rangos": [{"m3_desde": 0, "m3_hasta": None, "precio_m3": 1.0}],
        })
        igv_cfg = self.config.get("igv") or {}
        igv_activo = bool(igv_cfg.get("activo", False))
        igv_porcentaje = float(igv_cfg.get("porcentaje", 0.18))

        calc = calculos.calcular_recibo(
            medidores=doc.get("medidores", []),
            categoria=cat_obj,
            igv_activo=igv_activo,
            igv_porcentaje=igv_porcentaje,
            recibo_anterior=None,
        )

        doc.update({
            "consumo_total_m3": calc["consumo_total_m3"],
            "monto_agua": calc["monto_agua"],
            "monto_desague": calc["monto_desague"],
            "cargo_fijo": calc["cargo_fijo"],
            "deuda_anterior": calc["deuda_anterior"],
            "igv": calc["igv"],
            "redondeo": calc["redondeo"],
            "total_a_pagar": calc["total_a_pagar"],
            "meses_adeudados": calc["meses_adeudados"],
        })
        if "monto_pagado" not in doc or doc.get("monto_pagado") is None:
            doc["monto_pagado"] = 0.0
            doc["fecha_pago"] = None

        self.mes[id_usuario] = doc
        return doc

    def generar_pdf_recibo(self, periodo: str, id_usuario: str) -> str:
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError("Usuario no encontrado.")
        from servicios.generador_pdf import generar_recibo_pdf_para_usuario
        return str(generar_recibo_pdf_para_usuario(doc, self.config, periodo))

    def generar_zip_recibos(self, periodo: str, sector: Optional[str] = None) -> Optional[str]:
        from servicios.generador_pdf import generar_recibo_pdf_para_usuario, empaquetar_recibos_en_zip
        for uid, doc in self.mes.items():
            if sector and doc.get("sector") != sector:
                continue
            if doc.get("nro_recibo"):
                generar_recibo_pdf_para_usuario(doc, self.config, periodo)
        ruta_zip = empaquetar_recibos_en_zip(periodo, sector)
        return str(ruta_zip) if ruta_zip else None

    def registrar_pago(
        self,
        periodo: str,
        id_usuario: str,
        monto_pagado: float,
        fecha_pago: Optional[str] = None,
    ) -> Dict[str, Any]:
        import datetime
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError(f"Usuario {id_usuario} no encontrado en el mes {periodo}.")
        doc["monto_pagado"] = round(float(monto_pagado), 2)
        doc["fecha_pago"] = fecha_pago or datetime.date.today().isoformat()
        self.ops["escrituras"] += 1
        return doc

    def anular_pago(self, periodo: str, id_usuario: str) -> Dict[str, Any]:
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError(f"Usuario {id_usuario} no encontrado en el mes {periodo}.")
        doc["monto_pagado"] = 0.0
        doc["fecha_pago"] = None
        self.ops["escrituras"] += 1
        return doc

    def exportar_reporte_morosos_excel(
        self,
        periodo: Optional[str] = None,
        ruta_salida: Optional[str] = None,
    ) -> str:
        from pathlib import Path
        import rutas
        p = Path(ruta_salida) if ruta_salida else rutas.ruta_salida(f"reporte_morosos_{periodo or '2026-10'}.xlsx")
        p.parent.mkdir(parents=True, exist_ok=True)
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Lista de Corte"
            ws["A1"] = "Nro Conexión"
            ws["B1"] = "Titular"
            wb.save(str(p))
        except Exception:
            p.write_text("excel")
        return str(p)
