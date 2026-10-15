"""
servicios/datos_contingencia.py — Repositorio alterno para operar sin internet desde el respaldo Excel (Fase 6).

Implementa la misma interfaz que ServicioDatos, cargando la información del último
archivo Excel de respaldo, permitiendo búsquedas, consultas, reimpresión de PDFs y
encolando cobros y lecturas en cola_pendiente.json para su posterior sincronización.
"""
from __future__ import annotations
import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import calculos
import codigos
import respaldo_excel
import rutas
from servicios.cola_offline import encolar_cambio
from servicios.datos import ServicioDatos


def buscar_ultimo_respaldo_excel() -> Optional[Path]:
    """Busca el archivo de respaldo más reciente en la carpeta respaldos/."""
    dir_resp = rutas.ruta_respaldos()
    archivos = list(dir_resp.glob("RESPALDO_*.xlsx"))
    if not archivos:
        # Buscar en la raíz
        dir_raiz = dir_resp.parent / "respaldos"
        if dir_raiz.exists():
            archivos = list(dir_raiz.glob("RESPALDO_*.xlsx"))
    if not archivos:
        return None
    # Ordenar por fecha de modificación descendente
    archivos.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return archivos[0]


class ServicioDatosContingencia(ServicioDatos):
    """Servicio de datos que opera 100% fuera de línea desde un respaldo Excel."""

    def __init__(self, ruta_archivo_excel: Optional[Path | str] = None) -> None:
        self.ruta_excel: Optional[Path] = Path(ruta_archivo_excel) if ruta_archivo_excel else buscar_ultimo_respaldo_excel()
        if not self.ruta_excel or not self.ruta_excel.exists():
            raise FileNotFoundError("No se encontró ningún archivo de respaldo Excel para el modo contingencia.")

        # Cargar datos desde el Excel
        datos = respaldo_excel.leer_respaldo_excel(self.ruta_excel)
        self.config: Dict[str, Any] = datos.get("config") or {}
        self.usuarios: Dict[str, Dict[str, Any]] = datos.get("usuarios") or {}
        self.mes: Dict[str, Dict[str, Any]] = datos.get("mes_docs") or {}
        self.periodo: str = datos.get("periodo") or "2026-10"
        self.es_contingencia: bool = True
        self.ops = {"lecturas": 0, "escrituras": 0}

    def cargar_config(self) -> Dict[str, Any]:
        return self.config

    def guardar_config(self, datos: Dict[str, Any]) -> None:
        self.config = dict(datos)

    def actualizar_config_campo(self, campo: str, valor: Any) -> None:
        partes = campo.split(".")
        nodo = self.config
        for p in partes[:-1]:
            nodo = nodo.setdefault(p, {})
        nodo[partes[-1]] = valor

    def cargar_usuarios(self, forzar: bool = False) -> Dict[str, Dict[str, Any]]:
        return self.usuarios

    def obtener_usuario(self, id_usuario: str) -> Optional[Dict[str, Any]]:
        return self.usuarios.get(id_usuario)

    def guardar_usuario(
        self, id_usuario: str, datos: Dict[str, Any], agregar_al_mes: bool = False
    ) -> None:
        d = dict(datos)
        d["_id"] = id_usuario
        self.usuarios[id_usuario] = d

    def actualizar_usuario(self, id_usuario: str, cambios: Dict[str, Any]) -> None:
        if id_usuario in self.usuarios:
            self.usuarios[id_usuario].update(cambios)

    def corregir_dni_usuario(self, id_viejo: str, id_nuevo: str) -> Dict[str, Any]:
        if id_viejo not in self.usuarios:
            raise ValueError(f"Usuario {id_viejo} no existe.")
        u = self.usuarios.pop(id_viejo)
        u["_id"] = id_nuevo
        u["id_usuario"] = id_nuevo
        self.usuarios[id_nuevo] = u
        return u

    def obtener_siguiente_conexion_libre(self) -> str:
        con_max = 0
        for u in self.usuarios.values():
            c = str(u.get("nro_conexion") or "")
            if c.isdigit():
                con_max = max(con_max, int(c))
        return f"{con_max + 1:06d}"

    def obtener_siguiente_id_sd(self) -> str:
        sd_max = 0
        for uid in self.usuarios.keys():
            if str(uid).startswith("SD-"):
                num = str(uid)[3:]
                if num.isdigit():
                    sd_max = max(sd_max, int(num))
        return f"SD-{sd_max + 1:03d}"

    def validar_medidores_unicos(
        self, medidores: List[str], id_usuario_actual: Optional[str] = None
    ) -> Tuple[bool, Optional[str]]:
        return True, None

    def cargar_mes(self, periodo: str) -> Dict[str, Dict[str, Any]]:
        return self.mes

    def verificar_puede_abrir_mes(self) -> Tuple[bool, str, str]:
        return False, "No se puede abrir un nuevo mes en Modo Contingencia.", ""

    def verificar_periodo_disponible(self, periodo: str) -> Tuple[bool, str]:
        return False, "Modo contingencia activo: solo lectura y cobros locales."

    def abrir_mes(self, *args, **kwargs) -> Dict[str, Any]:
        raise RuntimeError("No se permite abrir mes en modo contingencia (requiere conexión con el servidor).")

    def cerrar_lecturas(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        return {"periodo": self.periodo, "estado": "CERRADO"}

    def reabrir_lecturas(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        return {"periodo": self.periodo, "estado": "ABIERTO"}

    def guardar_lectura_medidor(
        self,
        periodo: str,
        id_usuario: str,
        nro_medidor: str,
        lectura_actual: int,
        novedad: str = "NORMAL",
        nota: str = "",
        registrado_por: str = "ADMIN",
    ) -> Dict[str, Any]:
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError(f"Usuario {id_usuario} no encontrado en el mes {periodo}.")

        med_obj = next((m for m in doc.get("medidores", []) if m.get("nro_medidor") == nro_medidor), None)
        if med_obj:
            med_obj["lectura_actual"] = int(lectura_actual)
            med_obj["fecha_lectura"] = datetime.date.today().isoformat()
            med_obj["estado_lectura"] = "LEIDO"
            med_obj["novedad"] = novedad or "NORMAL"
            med_obj["nota"] = nota or ""
            med_obj["registrado_por"] = registrado_por

        # Encolar para sincronización posterior
        encolar_cambio("lectura", {
            "periodo": periodo,
            "id_usuario": id_usuario,
            "nro_medidor": nro_medidor,
            "lectura_actual": lectura_actual,
            "novedad": novedad,
            "nota": nota,
            "registrado_por": registrado_por,
        })
        self.ops["escrituras"] += 1
        return {"id_usuario": id_usuario, "nro_medidor": nro_medidor}

    def verificar_estado_generacion_recibos(self, periodo: Optional[str] = None) -> Dict[str, Any]:
        rec_gen = sum(1 for d in self.mes.values() if d.get("nro_recibo"))
        tot_fac = sum(float(d.get("total_a_pagar") or 0.0) for d in self.mes.values())
        pagados = sum(1 for d in self.mes.values() if float(d.get("monto_pagado") or 0) >= float(d.get("total_a_pagar") or 0) and float(d.get("total_a_pagar") or 0) > 0)
        return {
            "periodo": self.periodo,
            "estado": "FACTURADO",
            "puede_generar": False,
            "faltantes": [],
            "total_usuarios": len(self.mes),
            "recibos_generados": rec_gen,
            "total_facturado": round(tot_fac, 2),
            "pagados": pagados,
            "pendientes": len(self.mes) - pagados,
        }

    def estimar_lectura_medidor(self, *args, **kwargs) -> Dict[str, Any]:
        return {}

    def estimar_todas_lecturas_faltantes(self, *args, **kwargs) -> Dict[str, Any]:
        return {"estimados": 0, "requieren_manual": []}

    def generar_recibos_mes(self, *args, **kwargs) -> Dict[str, Any]:
        raise RuntimeError("No se permite generar nuevos recibos masivos en modo contingencia.")

    def generar_recibo_usuario(self, periodo: str, id_usuario: str) -> Dict[str, Any]:
        import calculos
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError(f"Usuario {id_usuario} no encontrado en el mes {periodo}.")

        for m in doc.get("medidores", []):
            if m.get("lectura_actual") is None:
                m["lectura_actual"] = m.get("lectura_anterior") or 0
                m["estado_lectura"] = "ESTIMADO"

        ultimo_recibo_num = int(self.config.get("ultimo_recibo") or 0)
        if not doc.get("nro_recibo"):
            ultimo_recibo_num += 1
            doc["nro_recibo"] = f"{ultimo_recibo_num:09d}"
            self.config["ultimo_recibo"] = ultimo_recibo_num

        categorias = self.config.get("categorias", {})
        cat_cod = doc.get("codigo_categoria", "A")
        cat_obj = categorias.get(cat_cod) or {"cargo_fijo": 1.5, "tarifa_desague": 0.0, "rangos": [{"m3_desde": 0, "m3_hasta": None, "precio_m3": 1.0}]}

        calc = calculos.calcular_recibo(
            medidores=doc.get("medidores", []),
            categoria=cat_obj,
            igv_activo=False,
            igv_porcentaje=0.18,
            recibo_anterior=None,
        )

        doc.update({
            "consumo_total_m3": calc["consumo_total_m3"],
            "monto_agua": calc["monto_agua"],
            "monto_desague": calc["monto_desague"],
            "cargo_fijo": calc["cargo_fijo"],
            "deuda_anterior": calc["deuda_anterior"],
            "igv": calc["igv"],
            "redondeo": calc["redondeo"],
            "total_a_pagar": calc["total_a_pagar"],
            "meses_adeudados": calc["meses_adeudados"],
        })
        if "monto_pagado" not in doc or doc.get("monto_pagado") is None:
            doc["monto_pagado"] = 0.0
            doc["fecha_pago"] = None

        self.mes[id_usuario] = doc
        return doc

    def generar_pdf_recibo(self, periodo: str, id_usuario: str) -> str:
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError("Usuario no encontrado.")
        from servicios.generador_pdf import generar_recibo_pdf_para_usuario
        return str(generar_recibo_pdf_para_usuario(doc, self.config, periodo))

    def generar_zip_recibos(self, periodo: str, sector: Optional[str] = None) -> Optional[str]:
        from servicios.generador_pdf import generar_recibo_pdf_para_usuario, empaquetar_recibos_en_zip
        for uid, doc in self.mes.items():
            if sector and doc.get("sector") != sector:
                continue
            if doc.get("nro_recibo"):
                generar_recibo_pdf_para_usuario(doc, self.config, periodo)
        ruta_zip = empaquetar_recibos_en_zip(periodo, sector)
        return str(ruta_zip) if ruta_zip else None

    def registrar_pago(
        self,
        periodo: str,
        id_usuario: str,
        monto_pagado: float,
        fecha_pago: Optional[str] = None,
    ) -> Dict[str, Any]:
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError(f"Usuario {id_usuario} no encontrado en el mes {periodo}.")
        fecha = fecha_pago or datetime.date.today().isoformat()
        monto = round(float(monto_pagado), 2)
        doc["monto_pagado"] = monto
        doc["fecha_pago"] = fecha

        # Encolar para sincronización
        encolar_cambio("pago", {
            "periodo": periodo,
            "id_usuario": id_usuario,
            "monto_pagado": monto,
            "fecha_pago": fecha,
        })
        self.ops["escrituras"] += 1
        return doc

    def anular_pago(self, periodo: str, id_usuario: str) -> Dict[str, Any]:
        doc = self.mes.get(id_usuario)
        if not doc:
            raise ValueError(f"Usuario {id_usuario} no encontrado en el mes {periodo}.")
        doc["monto_pagado"] = 0.0
        doc["fecha_pago"] = None

        # Encolar para sincronización
        encolar_cambio("anular_pago", {
            "periodo": periodo,
            "id_usuario": id_usuario,
        })
        self.ops["escrituras"] += 1
        return doc

    def exportar_reporte_morosos_excel(
        self,
        periodo: Optional[str] = None,
        ruta_salida: Optional[str] = None,
    ) -> str:
        # Puede usar el generador directo
        p = Path(ruta_salida) if ruta_salida else rutas.ruta_salida(f"reporte_morosos_{periodo or self.periodo}.xlsx")
        p.parent.mkdir(parents=True, exist_ok=True)
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Lista de Corte"
            ws["A1"] = "Nro Conexión"
            ws["B1"] = "Titular"
            wb.save(str(p))
        except Exception:
            p.write_text("excel")
        return str(p)

    def listar_lectores(self) -> List[Dict[str, Any]]:
        return []

    def crear_lector(self, *args, **kwargs) -> Dict[str, Any]:
        raise RuntimeError("No se pueden gestionar lectores en modo contingencia.")

    def actualizar_lector(self, *args, **kwargs) -> Dict[str, Any]:
        raise RuntimeError("No se pueden gestionar lectores en modo contingencia.")

    def eliminar_lector(self, *args, **kwargs) -> None:
        raise RuntimeError("No se pueden gestionar lectores en modo contingencia.")
