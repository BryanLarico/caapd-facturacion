"""
ui/componentes/dialogo_actualizar.py — Centro de Sincronización con Lectores Móviles (rediseñado).

Flujo en DOS pasos dentro del mismo diálogo:
  Paso 1 · Elegir acción:
      - "Sobreescribe Nube"  → Enviar mes a los lectores (Excel ➔ Nube).
      - "Descarga Espejo"    → Traer lecturas de los lectores (Nube ➔ Excel).
  Paso 2 · Pregunta "¿Está seguro?" con la explicación de lo que hace cada botón:
      - Sobreescribe Nube: exige escribir la palabra CONFIRMAR (Regla 5a) antes de habilitar el botón.
      - Descarga Espejo:   basta con pulsar "Sí, traer lecturas".
"""
from __future__ import annotations
import os
from typing import Any, Dict, List, Optional

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QStackedWidget,
    QVBoxLayout,
    QWidget,
)

import rutas


class DialogoActualizarBaseDatos(QDialog):
    """Diálogo modal para seleccionar y confirmar la sincronización espejo con los lectores móviles."""

    ACCION_LOCAL = "local"
    ACCION_SUBIR_FIREBASE = "subir_firebase"
    ACCION_DESCARGAR_FIREBASE = "descargar_firebase"
    ACCION_CANCELAR = "cancelar"

    PALABRA_CLAVE = "CONFIRMAR"

    # Toda la información visual y explicativa de cada botón vive aquí.
    INFO_ACCIONES: Dict[str, Dict[str, Any]] = {
        ACCION_SUBIR_FIREBASE: {
            "icono": "☁️",
            "titulo": "Enviar mes a los lectores",
            "direccion": "Excel ➔ Nube",
            "badge": "Sobreescribe Nube",
            "color": "#dc2626",
            "color_hover": "#b91c1c",
            "fondo": "#fef2f2",
            "borde": "#fecaca",
            "texto": "#991b1b",
            "resumen": "Sube el mes activo y el mes anterior a Firebase para que los lectores los vean en sus celulares.",
            "hace": [
                "Toma el <b>mes activo y el mes anterior</b> desde tus Excel de la carpeta 'Base de datos'.",
                "Sube cada mes a su colección en Firebase en lotes rápidos para que aparezcan en la app de los lectores.",
                "<b>Borra</b> en la nube los usuarios sobrantes y los meses más antiguos que esos dos.",
            ],
            "aviso": "Esto BORRARÁ todo lo que hay en Firebase y lo reemplazará con tus Excel actuales. "
                     "Las lecturas que los lectores hayan registrado y que aún NO hayas descargado se perderán.",
            "boton": "Sí, enviar a la nube",
        },
        ACCION_DESCARGAR_FIREBASE: {
            "icono": "📥",
            "titulo": "Traer lecturas de los lectores",
            "direccion": "Nube ➔ Excel",
            "badge": "Descarga Espejo",
            "color": "#0369a1",
            "color_hover": "#075985",
            "fondo": "#eff6ff",
            "borde": "#bfdbfe",
            "texto": "#1e40af",
            "resumen": "Descarga las lecturas registradas por los lectores móviles y las guarda en tu Excel local.",
            "hace": [
                "Lee desde Firebase las lecturas que los lectores registraron en sus celulares.",
                "Las escribe en el Excel local del mes activo como una copia espejo.",
                "<b>No toca</b> cobros, dinero ni lecturas hechas por ADMIN (si hay conflicto, ADMIN manda).",
            ],
            "aviso": "Operación segura: solo agrega o actualiza lecturas de los lectores. "
                     "Tus cobros y lecturas ADMIN permanecen intactos.",
            "boton": "Sí, traer lecturas",
        },
    }

    def __init__(self, parent: Optional[QWidget] = None, servicio: Optional[Any] = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Sincronización con Lectores Móviles")
        self.setModal(True)
        self.setFixedWidth(680)
        self.setMinimumHeight(700)
        self.setStyleSheet(self._hoja_estilos())

        self.resultado_accion = self.ACCION_CANCELAR
        self.accion_pendiente: Optional[str] = None
        self.servicio = servicio or getattr(parent, "servicio", None)

        dir_bd = rutas.ruta_base_de_datos()
        self.archivos_excel: List = [
            p for p in dir_bd.glob("*.xlsx") if not p.name.startswith("~$") and p.is_file()
        ]

        self._construir_ui()

    # ------------------------------------------------------------------ #
    #  Estilos
    # ------------------------------------------------------------------ #
    @staticmethod
    def _hoja_estilos() -> str:
        return """
            QDialog { background-color: #ffffff; }
            QLabel { background: transparent; border: none; }

            /* Cabecera con degradado azul agua */
            QFrame#cabecera {
                background-color: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                                                  stop:0 #0369a1, stop:1 #0ea5e9);
                border-radius: 14px;
            }
            QFrame#cabecera QLabel { color: #ffffff; }
            QLabel#cabIcono { font-size: 30px; }
            QLabel#cabTitulo { font-size: 19px; font-weight: bold; }
            QLabel#cabSub { font-size: 12px; color: #e0f2fe; }
            QLabel#cabPaso {
                font-size: 11px; font-weight: bold; color: #0c4a6e;
                background-color: #e0f2fe; border-radius: 10px; padding: 4px 10px;
            }

            /* Chips informativos */
            QFrame#chip {
                background-color: #f8fafc;
                border: 1px solid #e2e8f0;
                border-radius: 12px;
            }
            QLabel#chipIcono { font-size: 18px; }
            QLabel#chipValor { font-size: 13px; font-weight: bold; color: #0f172a; }
            QLabel#chipEtiqueta { font-size: 10px; color: #64748b; }

            /* Tarjetas de acción */
            QFrame#tarjeta {
                background-color: #ffffff;
                border: 1px solid #e2e8f0;
                border-radius: 14px;
            }
            QFrame#tarjeta:hover {
                border: 2px solid #0284c7;
                background-color: #f0f9ff;
            }
            QLabel#tarjetaTitulo { font-size: 14px; font-weight: bold; color: #0f172a; }
            QLabel#tarjetaDireccion { font-size: 11px; font-weight: 600; color: #0369a1; }
            QLabel#tarjetaDesc { font-size: 12px; color: #475569; }
            QLabel#chevron { font-size: 26px; color: #94a3b8; font-weight: bold; }

            /* Paso 2 · confirmación */
            QLabel#confPregunta { font-size: 22px; font-weight: bold; color: #0f172a; }
            QLabel#confSub { font-size: 13px; color: #475569; }
            QFrame#panelHace {
                background-color: #f8fafc;
                border: 1px solid #e2e8f0;
                border-radius: 12px;
            }
            QLabel#panelTitulo { font-size: 11px; font-weight: bold; color: #64748b; }
            QLabel#panelTexto { font-size: 12px; color: #334155; }
            QLabel#lblError { font-size: 11px; font-weight: 600; color: #b91c1c; }
            QLabel#lblAyuda { font-size: 11px; color: #64748b; }
            QLineEdit#txtConfirmar {
                border: 2px solid #fca5a5;
                border-radius: 10px;
                padding: 9px 14px;
                font-size: 15px;
                font-weight: bold;
                letter-spacing: 3px;
                color: #991b1b;
                background-color: #ffffff;
                selection-background-color: #fecaca;
            }
            QLineEdit#txtConfirmar:focus { border: 2px solid #dc2626; background-color: #fff7f7; }

            QFrame#panelDiferencia {
                background-color: #ffffff;
                border: 1px dashed #cbd5e1;
                border-radius: 12px;
            }

            /* Botones */
            QPushButton#btnEnlace {
                background: none; border: none; color: #0284c7;
                font-size: 12px; font-weight: 600; text-decoration: underline;
                padding: 4px 2px;
            }
            QPushButton#btnEnlace:hover { color: #0369a1; }

            QPushButton#btnSecundario, QPushButton#btnVolver {
                background-color: #f1f5f9; color: #334155;
                padding: 8px 20px; border: 1px solid #cbd5e1; border-radius: 9px;
                font-size: 13px; font-weight: 600; min-height: 20px;
            }
            QPushButton#btnSecundario:hover, QPushButton#btnVolver:hover { background-color: #e2e8f0; }

            QPushButton#btnSi {
                color: #ffffff; border: none; border-radius: 9px;
                padding: 8px 24px; font-size: 13px; font-weight: bold; min-height: 20px;
            }
            QPushButton#btnSi:disabled {
                background-color: #cbd5e1; color: #f8fafc;
            }
        """

    # ------------------------------------------------------------------ #
    #  Construcción de la interfaz
    # ------------------------------------------------------------------ #
    def _construir_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(22, 20, 22, 18)
        layout.setSpacing(14)

        layout.addWidget(self._crear_cabecera())

        self.pila = QStackedWidget()
        self.pila.addWidget(self._crear_pagina_seleccion())     # índice 0
        self.pila.addWidget(self._crear_pagina_confirmacion())  # índice 1
        layout.addWidget(self.pila, stretch=1)

    def _crear_cabecera(self) -> QFrame:
        cab = QFrame()
        cab.setObjectName("cabecera")
        lay = QHBoxLayout(cab)
        lay.setContentsMargins(18, 14, 18, 14)
        lay.setSpacing(14)

        lbl_icono = QLabel("🔄")
        lbl_icono.setObjectName("cabIcono")
        lay.addWidget(lbl_icono)

        col = QVBoxLayout()
        col.setSpacing(2)
        lbl_titulo = QLabel("Sincronización con Lectores Móviles")
        lbl_titulo.setObjectName("cabTitulo")
        col.addWidget(lbl_titulo)
        lbl_sub = QLabel("Intercambio de padrones y lecturas entre tus Excel locales y la app celular")
        lbl_sub.setObjectName("cabSub")
        lbl_sub.setWordWrap(True)
        col.addWidget(lbl_sub)
        lay.addLayout(col, stretch=1)

        self.lbl_cab_paso = QLabel("PASO 1 DE 2")
        self.lbl_cab_paso.setObjectName("cabPaso")
        self.lbl_cab_paso.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lay.addWidget(self.lbl_cab_paso, alignment=Qt.AlignmentFlag.AlignTop)

        return cab

    # ---------- Página 0: elegir acción ----------
    def _crear_pagina_seleccion(self) -> QWidget:
        pagina = QWidget()
        lay = QVBoxLayout(pagina)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(14)

        # Chips informativos (archivos / último envío / última descarga)
        cfg = getattr(self.servicio, "config", {}) if self.servicio else {}
        ts_envio = cfg.get("ultimo_envio_lectores") or "No registrado"
        ts_descarga = cfg.get("ultima_descarga_lectores") or "No registrada"

        n_arch = len(self.archivos_excel)
        nombres_tooltip = "\n".join(p.name for p in self.archivos_excel) or "Carpeta vacía."

        fila_chips = QHBoxLayout()
        fila_chips.setSpacing(10)
        fila_chips.addWidget(self._crear_chip("📁", f"{n_arch} archivo(s) Excel", "Carpeta 'Base de datos'", nombres_tooltip))
        fila_chips.addWidget(self._crear_chip("⬆️", str(ts_envio), "Último envío a lectores"))
        fila_chips.addWidget(self._crear_chip("⬇️", str(ts_descarga), "Última descarga"))
        lay.addLayout(fila_chips)

        lbl_elige = QLabel("¿Qué deseas hacer?")
        lbl_elige.setStyleSheet("font-size: 13px; font-weight: bold; color: #334155; margin-top: 4px;")
        lay.addWidget(lbl_elige)

        lay.addWidget(self._crear_tarjeta_accion(self.ACCION_SUBIR_FIREBASE))
        lay.addWidget(self._crear_tarjeta_accion(self.ACCION_DESCARGAR_FIREBASE))

        lbl_tip = QLabel(
            "💡 Al elegir una opción se te preguntará <b>¿Está seguro?</b> y verás exactamente "
            "qué hará antes de ejecutar nada."
        )
        lbl_tip.setObjectName("lblAyuda")
        lbl_tip.setWordWrap(True)
        lbl_tip.setTextFormat(Qt.TextFormat.RichText)
        lay.addWidget(lbl_tip)

        lay.addStretch(1)

        # Pie: abrir carpeta + cancelar
        pie = QHBoxLayout()
        btn_carpeta = QPushButton("📂 Abrir carpeta 'Base de datos'")
        btn_carpeta.setObjectName("btnEnlace")
        btn_carpeta.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_carpeta.clicked.connect(self._abrir_carpeta_bd)
        pie.addWidget(btn_carpeta)
        pie.addStretch(1)

        btn_cancelar = QPushButton("Cancelar")
        btn_cancelar.setObjectName("btnSecundario")
        btn_cancelar.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_cancelar.clicked.connect(self.reject)
        pie.addWidget(btn_cancelar)
        lay.addLayout(pie)

        return pagina

    def _crear_chip(self, icono: str, valor: str, etiqueta: str, tooltip: str = "") -> QFrame:
        chip = QFrame()
        chip.setObjectName("chip")
        if tooltip:
            chip.setToolTip(tooltip)
        lay = QHBoxLayout(chip)
        lay.setContentsMargins(12, 8, 12, 8)
        lay.setSpacing(10)

        lbl_ic = QLabel(icono)
        lbl_ic.setObjectName("chipIcono")
        lay.addWidget(lbl_ic)

        col = QVBoxLayout()
        col.setSpacing(0)
        lbl_val = QLabel(valor)
        lbl_val.setObjectName("chipValor")
        lbl_val.setWordWrap(True)
        col.addWidget(lbl_val)
        lbl_et = QLabel(etiqueta)
        lbl_et.setObjectName("chipEtiqueta")
        col.addWidget(lbl_et)
        lay.addLayout(col, stretch=1)
        return chip

    def _crear_tarjeta_accion(self, accion: str) -> QFrame:
        info = self.INFO_ACCIONES[accion]

        tarjeta = QFrame()
        tarjeta.setObjectName("tarjeta")
        tarjeta.setCursor(Qt.CursorShape.PointingHandCursor)
        tarjeta.setToolTip(f"Haz clic para elegir: {info['titulo']}")

        lay = QHBoxLayout(tarjeta)
        lay.setContentsMargins(16, 14, 16, 14)
        lay.setSpacing(14)

        # Círculo con ícono
        lbl_icono = QLabel(info["icono"])
        lbl_icono.setFixedSize(48, 48)
        lbl_icono.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_icono.setStyleSheet(
            f"background-color: {info['fondo']}; border: 1px solid {info['borde']}; "
            f"border-radius: 24px; font-size: 20px;"
        )
        lay.addWidget(lbl_icono, alignment=Qt.AlignmentFlag.AlignTop)

        # Textos
        col = QVBoxLayout()
        col.setSpacing(4)

        fila_titulo = QHBoxLayout()
        fila_titulo.setSpacing(8)
        lbl_tit = QLabel(info["titulo"])
        lbl_tit.setObjectName("tarjetaTitulo")
        fila_titulo.addWidget(lbl_tit)
        lbl_dir = QLabel(info["direccion"])
        lbl_dir.setObjectName("tarjetaDireccion")
        fila_titulo.addWidget(lbl_dir)
        fila_titulo.addStretch(1)
        lbl_badge = QLabel(info["badge"])
        lbl_badge.setStyleSheet(
            f"background-color: {info['fondo']}; color: {info['texto']}; "
            f"border: 1px solid {info['borde']}; border-radius: 9px; "
            f"font-size: 10px; font-weight: bold; padding: 3px 9px;"
        )
        fila_titulo.addWidget(lbl_badge)
        col.addLayout(fila_titulo)

        lbl_desc = QLabel(info["resumen"])
        lbl_desc.setObjectName("tarjetaDesc")
        lbl_desc.setWordWrap(True)
        col.addWidget(lbl_desc)

        lay.addLayout(col, stretch=1)

        lbl_chevron = QLabel("›")
        lbl_chevron.setObjectName("chevron")
        lay.addWidget(lbl_chevron, alignment=Qt.AlignmentFlag.AlignVCenter)

        # Clic en toda la tarjeta
        tarjeta.mousePressEvent = lambda event, a=accion: self._mostrar_confirmacion(a)
        return tarjeta

    # ---------- Página 1: ¿Está seguro? ----------
    def _crear_pagina_confirmacion(self) -> QWidget:
        pagina = QWidget()
        lay = QVBoxLayout(pagina)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(12)

        # Encabezado de la pregunta
        fila = QHBoxLayout()
        fila.setSpacing(16)
        self.lbl_conf_icono = QLabel("☁️")
        self.lbl_conf_icono.setFixedSize(60, 60)
        self.lbl_conf_icono.setAlignment(Qt.AlignmentFlag.AlignCenter)
        fila.addWidget(self.lbl_conf_icono, alignment=Qt.AlignmentFlag.AlignTop)

        col = QVBoxLayout()
        col.setSpacing(4)
        lbl_pregunta = QLabel("¿Está seguro?")
        lbl_pregunta.setObjectName("confPregunta")
        col.addWidget(lbl_pregunta)
        self.lbl_conf_sub = QLabel("")
        self.lbl_conf_sub.setObjectName("confSub")
        self.lbl_conf_sub.setWordWrap(True)
        col.addWidget(self.lbl_conf_sub)
        fila.addLayout(col, stretch=1)
        lay.addLayout(fila)

        # ¿Qué hará este botón?
        panel_hace = QFrame()
        panel_hace.setObjectName("panelHace")
        lay_h = QVBoxLayout(panel_hace)
        lay_h.setContentsMargins(14, 10, 14, 10)
        lay_h.setSpacing(6)
        lbl_h_tit = QLabel("¿QUÉ HARÁ ESTE BOTÓN?")
        lbl_h_tit.setObjectName("panelTitulo")
        lay_h.addWidget(lbl_h_tit)
        self.lbl_conf_hace = QLabel("")
        self.lbl_conf_hace.setObjectName("panelTexto")
        self.lbl_conf_hace.setWordWrap(True)
        self.lbl_conf_hace.setTextFormat(Qt.TextFormat.RichText)
        lay_h.addWidget(self.lbl_conf_hace)
        lay.addWidget(panel_hace)

        # Aviso (rojo o azul según la acción)
        self.panel_aviso = QFrame()
        self.panel_aviso.setObjectName("panelAviso")
        lay_a = QHBoxLayout(self.panel_aviso)
        lay_a.setContentsMargins(14, 10, 14, 10)
        lay_a.setSpacing(10)
        self.lbl_aviso_icono = QLabel("⚠️")
        self.lbl_aviso_icono.setStyleSheet("font-size: 18px;")
        lay_a.addWidget(self.lbl_aviso_icono, alignment=Qt.AlignmentFlag.AlignTop)
        self.lbl_aviso = QLabel("")
        self.lbl_aviso.setWordWrap(True)
        lay_a.addWidget(self.lbl_aviso, stretch=1)
        lay.addWidget(self.panel_aviso)

        # Caja CONFIRMAR (solo para Sobreescribe Nube)
        self.caja_confirmar = QWidget()
        lay_c = QVBoxLayout(self.caja_confirmar)
        lay_c.setContentsMargins(0, 0, 0, 0)
        lay_c.setSpacing(5)
        lbl_ayuda = QLabel(
            f"Para habilitar el botón, escribe exactamente la palabra <b>{self.PALABRA_CLAVE}</b> (en mayúsculas):"
        )
        lbl_ayuda.setObjectName("lblAyuda")
        lbl_ayuda.setWordWrap(True)
        lay_c.addWidget(lbl_ayuda)
        self.txt_confirmar = QLineEdit()
        self.txt_confirmar.setObjectName("txtConfirmar")
        self.txt_confirmar.setPlaceholderText(self.PALABRA_CLAVE)
        self.txt_confirmar.setMaxLength(20)
        self.txt_confirmar.textChanged.connect(self._al_escribir_confirmacion)
        self.txt_confirmar.returnPressed.connect(self._confirmar_accion)
        lay_c.addWidget(self.txt_confirmar)
        self.lbl_error_confirmar = QLabel("")
        self.lbl_error_confirmar.setObjectName("lblError")
        self.lbl_error_confirmar.setWordWrap(True)
        self.lbl_error_confirmar.hide()
        lay_c.addWidget(self.lbl_error_confirmar)
        lay.addWidget(self.caja_confirmar)

        # Recordatorio: qué hace cada botón
        panel_dif = QFrame()
        panel_dif.setObjectName("panelDiferencia")
        lay_d = QVBoxLayout(panel_dif)
        lay_d.setContentsMargins(14, 10, 14, 10)
        lay_d.setSpacing(6)
        lbl_d_tit = QLabel("RECUERDA LA DIFERENCIA ENTRE LOS DOS BOTONES")
        lbl_d_tit.setObjectName("panelTitulo")
        lay_d.addWidget(lbl_d_tit)
        for accion in (self.ACCION_SUBIR_FIREBASE, self.ACCION_DESCARGAR_FIREBASE):
            info = self.INFO_ACCIONES[accion]
            lbl = QLabel(
                f"{info['icono']}&nbsp; <b style='color:{info['texto']};'>{info['badge']}</b> "
                f"<span style='color:#64748b;'>({info['direccion']})</span> — {info['resumen']}"
            )
            lbl.setObjectName("panelTexto")
            lbl.setWordWrap(True)
            lbl.setTextFormat(Qt.TextFormat.RichText)
            lay_d.addWidget(lbl)
        lay.addWidget(panel_dif)

        lay.addStretch(1)

        # Pie: Volver / Sí
        pie = QHBoxLayout()
        btn_volver = QPushButton("‹  Volver")
        btn_volver.setObjectName("btnVolver")
        btn_volver.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_volver.clicked.connect(self._volver_a_seleccion)
        pie.addWidget(btn_volver)
        pie.addStretch(1)
        self.btn_si = QPushButton("Sí, continuar")
        self.btn_si.setObjectName("btnSi")
        self.btn_si.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_si.clicked.connect(self._confirmar_accion)
        pie.addWidget(self.btn_si)
        lay.addLayout(pie)

        return pagina

    # ------------------------------------------------------------------ #
    #  Navegación entre pasos
    # ------------------------------------------------------------------ #
    def _mostrar_confirmacion(self, accion: str) -> None:
        """Muestra la pregunta '¿Está seguro?' para la acción elegida."""
        if accion == self.ACCION_SUBIR_FIREBASE and self._hay_archivo_de_prueba():
            self._avisar_archivo_de_prueba()
            return

        self.accion_pendiente = accion
        info = self.INFO_ACCIONES[accion]
        es_subir = accion == self.ACCION_SUBIR_FIREBASE

        self.lbl_cab_paso.setText("PASO 2 DE 2")

        self.lbl_conf_icono.setText(info["icono"])
        self.lbl_conf_icono.setStyleSheet(
            f"background-color: {info['fondo']}; border: 2px solid {info['borde']}; "
            f"border-radius: 30px; font-size: 26px;"
        )
        self.lbl_conf_sub.setText(
            f"Estás a punto de ejecutar <b>{info['titulo']}</b> "
            f"<span style='color:{info['texto']}; font-weight:bold;'>({info['badge']} · {info['direccion']})</span>."
        )
        self.lbl_conf_hace.setText("<br>".join(f"•&nbsp; {linea}" for linea in info["hace"]))

        self.panel_aviso.setStyleSheet(
            f"QFrame#panelAviso {{ background-color: {info['fondo']}; border: 1px solid {info['borde']}; border-radius: 12px; }}"
        )
        self.lbl_aviso_icono.setText("⚠️" if es_subir else "🛡️")
        self.lbl_aviso.setText(info["aviso"])
        self.lbl_aviso.setStyleSheet(f"font-size: 12px; font-weight: 600; color: {info['texto']};")

        self.caja_confirmar.setVisible(es_subir)
        self.txt_confirmar.clear()
        self.lbl_error_confirmar.hide()

        self.btn_si.setText(info["boton"])
        self.btn_si.setEnabled(not es_subir)
        self.btn_si.setStyleSheet(
            f"QPushButton#btnSi {{ background-color: {info['color']}; }}"
            f"QPushButton#btnSi:hover {{ background-color: {info['color_hover']}; }}"
            f"QPushButton#btnSi:disabled {{ background-color: #cbd5e1; color: #f8fafc; }}"
        )

        self.pila.setCurrentIndex(1)
        if es_subir:
            self.txt_confirmar.setFocus()
        else:
            self.btn_si.setFocus()

    def _volver_a_seleccion(self) -> None:
        self.accion_pendiente = None
        self.lbl_cab_paso.setText("PASO 1 DE 2")
        self.pila.setCurrentIndex(0)

    def _al_escribir_confirmacion(self, texto: str) -> None:
        self.btn_si.setEnabled(texto.strip() == self.PALABRA_CLAVE)
        if self.lbl_error_confirmar.isVisible():
            self.lbl_error_confirmar.hide()

    def _confirmar_accion(self) -> None:
        if self.accion_pendiente == self.ACCION_SUBIR_FIREBASE:
            self._elegir_subir()
        elif self.accion_pendiente == self.ACCION_DESCARGAR_FIREBASE:
            self._elegir_descargar()

    # ------------------------------------------------------------------ #
    #  Acciones finales
    # ------------------------------------------------------------------ #
    def _elegir_subir(self) -> bool:
        """Confirma el envío a Firebase. Exige la palabra CONFIRMAR (Regla 5a). Devuelve True si procede."""
        texto = self.txt_confirmar.text().strip()
        if texto != self.PALABRA_CLAVE:
            self.lbl_error_confirmar.setText(
                f"✖ Debes escribir exactamente {self.PALABRA_CLAVE} (en mayúsculas). "
                "El envío a Firebase no se ha realizado."
            )
            self.lbl_error_confirmar.show()
            self.txt_confirmar.setFocus()
            self.txt_confirmar.selectAll()
            return False

        self.resultado_accion = self.ACCION_SUBIR_FIREBASE
        self.accept()
        return True

    def _elegir_descargar(self) -> None:
        self.resultado_accion = self.ACCION_DESCARGAR_FIREBASE
        self.accept()

    # ------------------------------------------------------------------ #
    #  Seguridad: archivos de prueba
    # ------------------------------------------------------------------ #
    def _hay_archivo_de_prueba(self) -> bool:
        """Devuelve True si algún Excel está marcado como 'TIPO: ARCHIVO DE PRUEBA' en su hoja LEEME."""
        for a in self.archivos_excel:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(str(a), read_only=True)
                try:
                    if "LEEME" in wb.sheetnames:
                        ws = wb["LEEME"]
                        for row in ws.iter_rows(max_row=20, max_col=2, values_only=True):
                            if row and row[0] and "ARCHIVO DE PRUEBA" in str(row[0]).upper():
                                return True
                finally:
                    wb.close()
            except Exception:
                pass
        return False

    def _avisar_archivo_de_prueba(self) -> None:
        msg = QMessageBox(self)
        msg.setIcon(QMessageBox.Icon.Critical)
        msg.setWindowTitle("Acceso Denegado por Seguridad")
        msg.setText("<b>Se detectó un archivo marcado como 'TIPO: ARCHIVO DE PRUEBA':</b>")
        msg.setInformativeText(
            "Por seguridad, los archivos de prueba NO se pueden subir al proyecto de producción real.\n\n"
            "Retire esa marca solo después de verificar que el archivo contiene datos reales."
        )
        msg.setStandardButtons(QMessageBox.StandardButton.Ok)
        msg.exec()

    # ------------------------------------------------------------------ #
    #  Utilidades
    # ------------------------------------------------------------------ #
    def _abrir_carpeta_bd(self) -> None:
        dir_bd = rutas.ruta_base_de_datos()
        try:
            os.startfile(str(dir_bd))
        except Exception:
            pass